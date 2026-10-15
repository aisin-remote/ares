"""Extract the authoritative supplier master rows from sheet SUPPLIER LIST."""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl


DAY_NUMBER = {
    "SENIN": 1,
    "SELASA": 2,
    "RABU": 3,
    "KAMIS": 4,
    "JUMAT": 5,
    "SABTU": 6,
    "MINGGU": 7,
}


def text(value: object) -> str:
    return " ".join(str(value).split()).strip()


def source_slug(value: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "-", value.upper()).strip("-")


def workdays(label: str) -> str:
    upper = label.upper()
    if upper == "DAILY":
        return "1,2,3,4,5"
    days = [DAY_NUMBER[item.strip()] for item in upper.split(",") if item.strip()]
    if not days:
        raise ValueError(f"DAY kosong/tidak valid: {label!r}")
    return ",".join(str(day) for day in sorted(set(days)))


def extract(source: Path) -> list[dict[str, str]]:
    sheet = openpyxl.load_workbook(source, data_only=True)["SUPPLIER LIST"]
    result: list[dict[str, str]] = []
    seen: dict[str, dict[str, str]] = {}
    for row in range(6, sheet.max_row + 1):
        plant = text(sheet.cell(row, 2).value) if sheet.cell(row, 2).value else ""
        supplier = text(sheet.cell(row, 4).value) if sheet.cell(row, 4).value else ""
        if not plant or not supplier:
            continue
        day_label = text(sheet.cell(row, 7).value)
        source_key = f"SUPPLIER-LIST:{source_slug(plant)}:{source_slug(supplier)}"
        item = {
            "plantCode": plant.upper(),
            "truckStation": text(sheet.cell(row, 3).value),
            "name": supplier,
            "scheduleCode": supplier.upper(),
            "deliveryCycle": text(sheet.cell(row, 5).value),
            "truckType": text(sheet.cell(row, 6).value),
            "deliveryDayLabel": day_label.upper(),
            "deliveryDays": workdays(day_label),
            "palletPerDelivery": text(sheet.cell(row, 8).value),
            "loadingOperator": text(sheet.cell(row, 9).value),
            "location": text(sheet.cell(row, 10).value),
            "sourceKey": source_key,
        }
        previous = seen.get(source_key)
        if previous and previous != item:
            raise ValueError(f"Supplier duplikat dengan isi berbeda: {source_key}")
        if previous:
            continue
        seen[source_key] = item
        result.append(item)
    return result


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("usage: extract-supplier-list.py SOURCE.xlsx OUTPUT.ts")
    source = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    rows = extract(source)
    output.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(rows, ensure_ascii=False, indent=2)
    output.write_text(
        "// Generated from sheet SUPPLIER LIST in Truck Station Diagram.xlsx.\n"
        "// Regenerate with scripts/extract-supplier-list.py; do not edit by hand.\n"
        f"export const REAL_SUPPLIERS = {payload} as const;\n",
        encoding="utf-8",
    )
    counts: dict[str, int] = {}
    for row in rows:
        counts[row["plantCode"]] = counts.get(row["plantCode"], 0) + 1
    print(f"wrote {len(rows)} suppliers {counts} to {output}")


if __name__ == "__main__":
    main()
